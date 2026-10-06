#!/usr/bin/env python
"""Gera dados_exemplo/Analise_Query.xlsx com dados 100% FICTÍCIOS.

O arquivo tem as abas e colunas que js/models/PlanilhaModel.js espera:
  Unidades (cadastro), Maquinas (PDV -> unidade), Cartões e Pix (transações).
Seed fixa: rodar de novo produz sempre o mesmo conteúdo.

Uso:  pip install openpyxl   &&   python gerar_dados_exemplo.py
"""
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

SEED = 20260101
N_UNIDADES = 180
N_MAQUINAS = 140
N_CARTOES = 9000
N_PIX = 3500
INICIO = date(2025, 1, 1)
DIAS = 180

# UFs reais do Brasil com a região correspondente (as unidades são inventadas).
UFS = {
    "SP": "SUDESTE", "MG": "SUDESTE", "RJ": "SUDESTE", "ES": "SUDESTE",
    "PR": "SUL", "RS": "SUL", "SC": "SUL",
    "BA": "NORDESTE", "CE": "NORDESTE", "PE": "NORDESTE", "MA": "NORDESTE", "RN": "NORDESTE",
    "PA": "NORTE", "AM": "NORTE", "TO": "NORTE",
    "GO": "CENTRO OESTE", "MT": "CENTRO OESTE", "MS": "CENTRO OESTE", "DF": "CENTRO OESTE",
}
PESO_UF = {"SP": 10, "MG": 7, "RJ": 4, "PR": 4, "RS": 4, "BA": 4, "PA": 3, "GO": 3}
DESCRICOES = ["UNIDADE LOCAL", "UNIDADE REGIONAL", "UNIDADE CENTRAL", "UNIDADE SETORIAL",
              "POSTO DE APOIO", "NÚCLEO COMUNITÁRIO"]
PESO_DESC = [8, 3, 1, 1, 1, 1]
IMOVEL = ["Próprio", "Alugado", "Cedido"]
MODELOS = ["MODELO A100", "MODELO A100 PRO", "MODELO B200", "MODELO C300"]
BANDEIRAS = ["Mastercard", "Visa", "Elo", "American Express"]
PESO_BAND = [44, 37, 18, 1]


def main():
    rng = random.Random(SEED)
    ufs = list(UFS)
    pesos = [PESO_UF.get(u, 1) for u in ufs]

    wb = Workbook(write_only=True)

    # ---- Unidades
    ws = wb.create_sheet("Unidades")
    ws.append(["ID", "Nome", "N Fantasia", "Endereco", "Bairro", "Estado", "Municipio",
               "Desc.Região", "Descrição da Unidade", "Status", "Qtd.Membros", "Tipo Imovel"])
    unidades = []
    for i in range(1, N_UNIDADES + 1):
        uf = rng.choices(ufs, pesos)[0]
        cidade = f"Cidade {rng.randint(1, 12):02d}"
        desc = rng.choices(DESCRICOES, PESO_DESC)[0]
        unidades.append(i)
        ws.append([1000 + i, f"Unidade {i:03d}", f"Unidade {i:03d} Exemplo", f"Rua Exemplo, {rng.randint(1, 999)}",
                   f"Bairro {rng.randint(1, 20):02d}", uf, cidade, UFS[uf], desc,
                   rng.choices(["Ativo", "Inativo"], [7, 3])[0], rng.randint(20, 900), rng.choice(IMOVEL)])

    # ---- Maquinas (alguns PDVs ficam sem unidade para exercitar "Sem localidade")
    ws = wb.create_sheet("Maquinas")
    ws.append(["PDV", "Nº de série", "Modelo", "ID_Unidade"])
    pdvs = []
    donas = rng.sample(unidades, N_MAQUINAS)
    for i in range(N_MAQUINAS):
        pdv = 50000 + i
        pdvs.append(pdv)
        uid = 1000 + donas[i] if i % 40 != 39 else None
        ws.append([pdv, f"SN{rng.randint(10**7, 10**8 - 1)}", rng.choice(MODELOS), uid])

    # ---- transações: cada PDV tem um "porte" (peso) que dá receitas bem diferentes
    porte = {p: rng.paretovariate(1.6) for p in pdvs}
    pesos_pdv = [porte[p] for p in pdvs]

    def data():
        d = INICIO + timedelta(days=rng.randrange(DIAS))
        return d

    def status():
        return rng.choices(["Aprovada", "Negada", "Cancelada"], [92, 5, 3])[0]

    ws = wb.create_sheet("Cartões")
    ws.append(["Data", "Pdv", "Tipo", "Bandeira", "Valor", "Status"])
    for _ in range(N_CARTOES):
        tipo = rng.choices(["Débito", "Crédito"], [75, 25])[0]
        base = rng.lognormvariate(3.3, 0.9) * (1.6 if tipo == "Crédito" else 1)
        ws.append([data(), rng.choices(pdvs, pesos_pdv)[0], tipo,
                   rng.choices(BANDEIRAS, PESO_BAND)[0], round(base, 2), status()])

    ws = wb.create_sheet("Pix")
    ws.append(["Data", "Status", "PDV", "Valor"])
    for _ in range(N_PIX):
        ws.append([data(), status(), rng.choices(pdvs, pesos_pdv)[0],
                   round(rng.lognormvariate(3.1, 0.9), 2)])

    destino = Path(__file__).parent / "dados_exemplo"
    destino.mkdir(exist_ok=True)
    arquivo = destino / "Analise_Query.xlsx"
    wb.save(arquivo)
    print(f"gerado: {arquivo} ({arquivo.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
